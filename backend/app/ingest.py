"""Text extraction + chunking.

Kept deliberately dependency-light: pypdf for PDFs, python-docx for Word, and
plain decoding for everything text-shaped.
"""

from __future__ import annotations

import io
import re

from pypdf import PdfReader

from .config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx", ".csv", ".json"}


class UnsupportedFileError(ValueError):
    pass


def extract_text(filename: str, raw: bytes) -> str:
    """Return the plain text of an uploaded file."""
    lower = filename.lower()

    if lower.endswith(".pdf"):
        return _extract_pdf(raw)
    if lower.endswith(".docx"):
        return _extract_docx(raw)
    if any(lower.endswith(ext) for ext in (".txt", ".md", ".markdown", ".csv", ".json")):
        return raw.decode("utf-8", errors="replace")

    raise UnsupportedFileError(
        f"Unsupported file type. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _extract_pdf(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = _normalise(text)
        if text:
            # Page markers survive chunking and make citations far more useful.
            pages.append(f"[page {i}]\n{text}")
    return "\n\n".join(pages)


def _extract_docx(raw: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _normalise("\n".join(parts))


def _normalise(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(
    text: str,
    chunk_size: int | None = None,
    overlap: int | None = None,
) -> list[str]:
    """Split text into overlapping chunks, preferring paragraph then sentence
    boundaries so a chunk rarely stops mid-thought."""
    chunk_size = chunk_size or settings.chunk_size
    overlap = overlap or settings.chunk_overlap
    if overlap >= chunk_size:
        overlap = chunk_size // 5

    text = _normalise(text)
    if not text:
        return []

    # Build up chunks from paragraphs; split a paragraph further only when it
    # is itself larger than the target chunk size.
    units: list[str] = []
    for para in text.split("\n\n"):
        para = para.strip()
        if not para:
            continue
        if len(para) <= chunk_size:
            units.append(para)
        else:
            units.extend(_split_long(para, chunk_size))

    chunks: list[str] = []
    current = ""
    for unit in units:
        candidate = f"{current}\n\n{unit}" if current else unit
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = f"{tail}\n\n{unit}".strip() if tail else unit
        else:
            current = unit

    if current.strip():
        chunks.append(current.strip())

    return [c for c in (c.strip() for c in chunks) if c]


def _split_long(para: str, chunk_size: int) -> list[str]:
    """Break an oversized paragraph on sentence boundaries, falling back to a
    hard character cut for pathological input (e.g. a wall of base64)."""
    sentences = re.split(r"(?<=[.!?])\s+", para)
    out: list[str] = []
    current = ""
    for sentence in sentences:
        while len(sentence) > chunk_size:
            out.append(sentence[:chunk_size])
            sentence = sentence[chunk_size:]
        candidate = f"{current} {sentence}".strip()
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current:
                out.append(current)
            current = sentence
    if current:
        out.append(current)
    return out
